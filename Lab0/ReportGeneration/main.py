import sys
import json
from openpyxl import Workbook
from openpyxl.styles import Border, Side
from copy import copy

thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))


def add_table(wb, table):
    name = table['Name']
    wb.create_sheet(name)
    ws = wb[name]
    records = table['Records']
    header = [attribute for attribute in records[0]]
    ws.append(header)
    for record in records:
        ws.append(value for value in record.values())


def decorate_workbook(wb):
    for ws in wb.worksheets:
        for column in ws.columns:
            column_identifier = column[0].column_letter
            max_width = 0
            for cell in column:
                current_width = len(str(cell.value))
                if current_width > max_width:
                    max_width = current_width
                cell.alignment = get_alignment_object(cell)
                cell.border = thin_border
            adjusted_width = max_width * 1.25
            ws.column_dimensions[column_identifier].width = adjusted_width


def get_alignment_object(cell):
    alignment_object = copy(cell.alignment)
    alignment_object.horizontal = 'center'
    alignment_object.vertical = 'center'
    return alignment_object


def main():
    json_path = sys.argv[1]
    with open(json_path, "r") as file:
        tables = json.load(file)

    wb = Workbook()
    del wb['Sheet']
    for table in tables:
        add_table(wb, table)

    decorate_workbook(wb)

    excel_path = sys.argv[2]
    wb.save(excel_path)
    print('Done')


if __name__ == '__main__':
    main()
